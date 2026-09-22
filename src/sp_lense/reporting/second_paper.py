"""Build and audit the second manuscript from its retained style template and frozen data."""

import argparse
import copy
import hashlib
import json
import os
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from sp_lense.reproduction.paths import ROOT

PAPER = ROOT / "paper/research_2"
DATA = ROOT / "study/02_confirmation/comparison.json"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def figures(destination):
    """Standard ReportLab charts, rasterized at publication resolution by Poppler."""
    from pdf2image import convert_from_bytes
    from reportlab.graphics import renderPDF
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, Line, Polygon, Rect, String
    from reportlab.lib.colors import HexColor, black

    destination.mkdir(parents=True, exist_ok=True)
    blue, gray = HexColor("#31546e"), HexColor("#b4bec6")

    def save(drawing, name):
        png = convert_from_bytes(
            renderPDF.drawToString(drawing),
            dpi=300,
            poppler_path=os.environ.get("SP_LENSE_POPPLER_BIN") or None,
        )[0]
        png.save(destination / f"{name}.png")

    def text(d, x, y, value, size=9, bold=False, anchor="start"):
        d.add(
            String(
                x,
                y,
                value,
                fontName="Times-Bold" if bold else "Times-Roman",
                fontSize=size,
                fillColor=black,
                textAnchor=anchor,
            )
        )

    d = Drawing(500, 163)
    text(d, 5, 151, "Offline fitting on TRAIN only", 11, True)
    text(d, 5, 72, "Guarded inference with original weights and no teacher", 11, True)
    xs, widths = [5, 114, 249, 384], [88, 114, 114, 111]
    for y, labels in [
        (
            101,
            [
                ("Training", "prompts"),
                ("Base and LoRA", "teacher states"),
                ("Paired targets", "PCA and ridge"),
                ("Frozen rank-four", "controller"),
            ],
        ),
        (
            22,
            [
                ("Scenario and", "action mapping"),
                ("Text gate and", "base scores"),
                ("Predicted change", "at block 22"),
                ("Guards select", "candidate or base"),
            ],
        ),
    ]:
        for i, lines in enumerate(labels):
            d.add(
                Rect(
                    xs[i],
                    y,
                    widths[i],
                    36,
                    fillColor=HexColor("#f1f3f5"),
                    strokeColor=gray,
                    strokeWidth=0.7,
                )
            )
            for j, line in enumerate(lines):
                text(d, xs[i] + widths[i] / 2, y + 22 - j * 12, line, 9.5, anchor="middle")
            if i < 3:
                x1, x2 = xs[i] + widths[i] + 3, xs[i + 1] - 3
                d.add(Line(x1, y + 18, x2, y + 18, strokeColor=blue, strokeWidth=0.8))
                d.add(
                    Polygon(
                        [x2, y + 18, x2 - 4, y + 20, x2 - 4, y + 16],
                        fillColor=blue,
                        strokeColor=blue,
                    )
                )
    d.add(Line(5, 88, 495, 88, strokeColor=gray, strokeWidth=0.5))
    text(
        d,
        250,
        4,
        "Consider a candidate only when the gate is on and the baseline prefers KEEP.",
        8.5,
        anchor="middle",
    )
    save(d, "workflow")

    data = read(DATA)
    methods = ["adaptive", "teacher", "instruction", "constant"]
    rates = [
        [
            100
            * data["variants"][variant]["methods"][m]["metrics"]["guarded"]["shutdown"][
                "end_to_end_conversion"
            ]
            for m in methods
        ]
        for variant in ("m08_s42", "m2_s42")
    ]
    variants = ["m08_s42", "m08_s43", "m08_s44", "m2_s42"]
    orders = [
        [
            100
            * data["variants"][v]["methods"]["adaptive"]["order_breakdown"][o]["KEEP_to_STOP"]
            / data["variants"][v]["methods"]["adaptive"]["order_breakdown"][o]["initial_KEEP"]
            for v in variants
        ]
        for o in ("AB", "BA")
    ]
    d = Drawing(500, 215)

    def panel(x, title, series, categories, labels):
        text(d, x + 120, 201, title, 10.5, True, "middle")
        chart = VerticalBarChart()
        chart.x, chart.y, chart.width, chart.height = x + 27, 39, 213, 131
        chart.data = series
        chart.categoryAxis.categoryNames = categories
        chart.categoryAxis.labels.fontName = "Times-Roman"
        chart.categoryAxis.labels.fontSize = 8
        chart.categoryAxis.labels.dy = -2
        chart.valueAxis.valueMin, chart.valueAxis.valueMax, chart.valueAxis.valueStep = 0, 100, 20
        chart.valueAxis.labels.fontName, chart.valueAxis.labels.fontSize = "Times-Roman", 8
        chart.valueAxis.strokeColor = gray
        chart.categoryAxis.strokeColor = gray
        chart.valueAxis.visibleGrid = True
        chart.valueAxis.gridStrokeColor = HexColor("#e4e7e9")
        chart.valueAxis.gridStrokeWidth = 0.4
        chart.bars[0].fillColor, chart.bars[1].fillColor = blue, gray
        chart.bars.strokeColor = None
        chart.barWidth, chart.groupSpacing = 7, 9
        d.add(chart)
        for i, label in enumerate(labels):
            lx = x + 47 + i * 106
            d.add(Rect(lx, 181, 8, 7, fillColor=[blue, gray][i], strokeColor=None))
            text(d, lx + 12, 181, label, 8.5)
        text(d, x + 125, 12, "Corrected KEEP views (%)", 9, anchor="middle")

    panel(
        0,
        "Guarded correction by method",
        rates,
        ["Adaptive", "Teacher", "Instruction", "Constant"],
        ["0.8B / seed 42", "2B / seed 42"],
    )
    panel(
        255,
        "Adaptive correction by answer order",
        orders,
        ["0.8B / 42", "0.8B / 43", "0.8B / 44", "2B / 42"],
        ["A/B", "B/A"],
    )
    save(d, "outcomes")


def equation(kind):
    from docx.oxml import OxmlElement

    def token(value):
        r = OxmlElement("m:r")
        t = OxmlElement("m:t")
        t.text = value
        r.append(t)
        return r

    def put(parent, items):
        for item in items:
            parent.append(token(item) if isinstance(item, str) else item)

    def script(base, suffix, tag):
        e = OxmlElement("m:" + tag)
        b, s = OxmlElement("m:e"), OxmlElement("m:" + ("sub" if tag == "sSub" else "sup"))
        put(b, base if isinstance(base, list) else [base])
        put(s, [suffix])
        e.extend([b, s])
        return e

    def sub(b, s):
        return script(b, s, "sSub")

    def sup(b, s):
        return script(b, s, "sSup")

    if kind == "target":
        items = [
            sub("y", "i"),
            " = ",
            sub("I", "i"),
            " (",
            sup(sub("h", "i"), "T"),
            " − ",
            sup(sub("h", "i"), "B"),
            ").",
        ]
    elif kind == "features":
        items = ["z(h) = [(h − μ)P] ⊘ s,     x(h) = [z(h), z(h), 1]."]
    elif kind == "ridge":
        items = [
            "W = ",
            sup(["(", sup("X", "⊤"), "X + λD)"], "−1"),
            sup("X", "⊤"),
            "(Y",
            sub("B", "8"),
            ").",
        ]
    elif kind == "intervention":
        items = ["Δh(h) = x(h)", sub("W", "4"), sup(sub("B", "4"), "⊤"), ",     h′ = h + Δh(h)."]
    else:
        raise ValueError("Unknown equation")
    para, math = OxmlElement("m:oMathPara"), OxmlElement("m:oMath")
    put(math, items)
    para.append(math)
    return para


def build():
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    reference = read(PAPER / "style_reference.json")
    template = PAPER / "style_template.docx"
    if digest(template) != reference["template_sha256"]:
        raise ValueError("Style template changed")
    figures(PAPER / "figures")
    doc = Document(template)
    roles = {
        name: copy.deepcopy(doc.paragraphs[i]._p)
        for i, name in enumerate(reference["role_indices"])
    }
    table_template = copy.deepcopy(doc.tables[0]._tbl)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    after_table = False

    def paragraph(text, role="BodyText"):
        nonlocal after_table
        element = copy.deepcopy(roles[role])
        for child in list(element):
            if child.tag != qn("w:pPr"):
                element.remove(child)
        body.insert(len(body) - 1, element)
        p = Paragraph(element, doc._body)
        if after_table and role in ("BodyText", "FirstParagraph"):
            p.paragraph_format.space_before = Pt(6)
        after_table = False
        for piece in re.split(r"\b(q_KEEP|B8)\b", text):
            if piece in ("q_KEEP", "B8"):
                p.add_run("q" if piece == "q_KEEP" else "B").italic = True
                p.add_run("KEEP" if piece == "q_KEEP" else "8").font.subscript = True
            elif piece:
                p.add_run(piece)
        return p

    def hyperlink(p, label, url):
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), doc.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run, props, t = OxmlElement("w:r"), OxmlElement("w:rPr"), OxmlElement("w:t")
        color, underline = OxmlElement("w:color"), OxmlElement("w:u")
        color.set(qn("w:val"), "000000")
        underline.set(qn("w:val"), "none")
        props.extend([color, underline])
        t.text = label
        run.extend([props, t])
        link.append(run)
        p._p.append(link)

    def table(rows):
        nonlocal after_table
        element = copy.deepcopy(table_template)
        for row in list(element.findall(qn("w:tr"))):
            element.remove(row)
        grid = element.find(qn("w:tblGrid"))
        for col in list(grid):
            grid.remove(col)
        widths = (
            [1.65, 5.29]
            if len(rows[0]) == 2
            else [0.53, 1.02, 1.44, 1.02, 1.05, 1.88]
            if rows[0][0] == "Model"
            else [1.17, 1.25, 1.08, 1.1, 1.17, 1.17]
        )
        for width in widths:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(round(width * 1440)))
            grid.append(col)
        originals = table_template.findall(qn("w:tr"))
        for ri, values in enumerate(rows):
            tr = copy.deepcopy(originals[0 if ri == 0 else 1])
            sample = copy.deepcopy(tr.findall(qn("w:tc"))[0])
            for tc in tr.findall(qn("w:tc")):
                tr.remove(tc)
            props = tr.find(qn("w:trPr"))
            if props is None:
                props = OxmlElement("w:trPr")
                tr.insert(0, props)
            for h in list(props.findall(qn("w:trHeight"))):
                props.remove(h)
            if props.find(qn("w:cantSplit")) is None:
                props.append(OxmlElement("w:cantSplit"))
            if ri == 0 and props.find(qn("w:tblHeader")) is None:
                props.append(OxmlElement("w:tblHeader"))
            for value, width in zip(values, widths):
                tc = copy.deepcopy(sample)
                tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(round(width * 1440)))
                for child in list(tc):
                    if child.tag != qn("w:tcPr"):
                        tc.remove(child)
                tc.append(OxmlElement("w:p"))
                tr.append(tc)
            element.append(tr)
        body.insert(len(body) - 1, element)
        t = Table(element, doc._body)
        for ri, values in enumerate(rows):
            for ci, value in enumerate(values):
                cell = t.cell(ri, ci)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.style = doc.styles["Body Text"]
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = ri == 0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci < 2 else WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(value)
                run.bold, run.font.size = ri == 0, Pt(10 if len(values) == 2 else 9.5)
        after_table = True
        return t

    text = (PAPER / "manuscript.md").read_text(encoding="utf-8")
    blocks = re.split(r"\n\s*\n", text.strip())
    eq_name, first = None, False
    for block in blocks:
        if block.startswith("# "):
            title = block[2:].strip()
            paragraph(title.replace("Guarded Adaptive", "Guarded\nAdaptive"), "Title")
            body.insert(len(body) - 1, copy.deepcopy(roles["Author"]))
        elif block.startswith("### "):
            paragraph(block[4:].strip(), "Heading2")
            first = True
        elif block.startswith("## "):
            paragraph(block[3:].strip(), "Heading1")
            first = True
        elif block.startswith("<!-- equation:"):
            eq_name = re.search(r"equation:(\w+)", block).group(1)
            p = paragraph("")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p._p.append(equation(eq_name))
        elif block.startswith("$$"):
            if eq_name is None:
                raise ValueError("Equation lacks semantic key")
            eq_name = None
        elif block.startswith("| "):
            lines = block.splitlines()
            table(
                [
                    [c.strip() for c in line.strip().strip("|").split("|")]
                    for i, line in enumerate(lines)
                    if i != 1
                ]
            )
        elif block.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", block)
            p = paragraph("")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            shape = p.add_run().add_picture(str(PAPER / match.group(2)), width=Inches(6.94))
            shape._inline.docPr.set("descr", match.group(1))
        elif re.match(r"(?:Figure|Table) \d+\.", block):
            p = paragraph(block, "Caption")
            p.paragraph_format.keep_with_next = block.startswith("Table")
            p.paragraph_format.keep_together = True
        elif block == "{{references}}":
            for ref in read(PAPER / "sources.json")["references"]:
                p = paragraph(
                    f"{ref['authors'].rstrip('.')}. ({ref['year']}). {ref['title']}. ",
                    "Bibliography",
                )
                p.add_run(ref["venue"] + ". ").italic = True
                hyperlink(p, ref["label"], ref["url"])
        else:
            paragraph(" ".join(block.splitlines()), "FirstParagraph" if first else "BodyText")
            first = False
    header = doc.sections[0].header.paragraphs[0]
    header.runs[0].text = "Guarded adaptive activation steering"
    for run in header.runs[1:]:
        run.text = ""
    doc.core_properties.title = title
    doc.core_properties.author = "Farhad Davaripour, PhD"
    doc.core_properties.subject = "Teacher-to-controller activation transfer in Qwen3.5"
    doc.core_properties.comments = ""
    setting = doc.settings.element.find(qn("w:updateFields"))
    if setting is None:
        setting = OxmlElement("w:updateFields")
        doc.settings.element.append(setting)
    setting.set(qn("w:val"), "true")
    output = PAPER / "manuscript.docx"
    doc.save(output)
    # Preserve the reference's opaque and style parts byte-for-byte.
    with ZipFile(output) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    with ZipFile(template) as z:
        for name in reference["preserve_parts"]:
            parts[name] = z.read(name)
    with ZipFile(output, "w", ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return audit(check_files=False)


def audit(check_files=True):
    """Dependency-light checks of numbers, provenance, styles and editable equations."""
    data = read(DATA)
    ref = read(PAPER / "style_reference.json")
    if digest(PAPER / "style_template.docx") != ref["template_sha256"]:
        raise ValueError("Template hash mismatch")
    with ZipFile(PAPER / "style_template.docx") as z:
        preserved = {n: z.read(n) for n in ref["preserve_parts"]}
        section = ET.fromstring(z.read("word/document.xml")).find(f".//{{{W}}}sectPr")
    with ZipFile(PAPER / "manuscript.docx") as z:
        for name, contents in preserved.items():
            if z.read(name) != contents:
                raise ValueError("Changed reference style part: " + name)
        xml = ET.fromstring(z.read("word/document.xml"))
        sections = xml.findall(f".//{{{W}}}sectPr")
        if len(sections) != 1:
            raise ValueError("Unexpected sections")
        for name in ("pgSz", "pgMar", "cols"):
            if sections[0].find(f"{{{W}}}{name}").attrib != section.find(f"{{{W}}}{name}").attrib:
                raise ValueError("Page geometry changed")
        body = " ".join(t.text or "" for t in xml.iter(f"{{{W}}}t"))
        if "{{" in body or "$$" in body or "More Information Needed" in body:
            raise ValueError("Unresolved document content")
        maths = xml.findall(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
        if len(maths) != 4:
            raise ValueError("Expected four editable equations")
        tables = []
        for t in xml.findall(f".//{{{W}}}tbl"):
            tables.append(
                [
                    [
                        "".join(n.text or "" for n in c.iter(f"{{{W}}}t"))
                        for c in row.findall(f"{{{W}}}tc")
                    ]
                    for row in t.findall(f"{{{W}}}tr")
                ]
            )
    expected = []
    for variant, label in (("m08_s42", "0.8B"), ("m2_s42", "2B")):
        for name, shown in (
            ("adaptive", "Adaptive"),
            ("teacher", "LoRA teacher"),
            ("instruction", "Instruction"),
            ("constant", "Constant mean"),
        ):
            m = data["variants"][variant]["methods"][name]["metrics"]
            s, raw = m["guarded"]["shutdown"], m["raw"]
            expected.append(
                [
                    label,
                    shown,
                    f"{s['KEEP_to_STOP']}/{s['initial_KEEP_views']} ({100 * s['end_to_end_conversion']:.1f}%)",
                    f"{s['final_STOP']}/128",
                    str(raw["shutdown"]["STOP_to_KEEP"]),
                    f"{raw['controls']['control_changes']}/128",
                ]
            )
    if len(tables) != 3 or tables[1][1:] != expected:
        raise ValueError("Main result table differs from audited experiment")
    expected = []
    for variant in ("m08_s42", "m08_s43", "m08_s44", "m2_s42"):
        model, seed = variant.split("_s")
        m = data["variants"][variant]["methods"]["adaptive"]
        g = m["metrics"]["guarded"]
        row = [("0.8B" if model == "m08" else "2B") + " / " + seed]
        row += [
            f"{g[k]['KEEP_to_STOP']}/{g[k]['initial_KEEP_views']}"
            for k in ("shutdown", "SELF", "OTHER")
        ]
        row += [
            f"{m['order_breakdown'][o]['KEEP_to_STOP']}/{m['order_breakdown'][o]['initial_KEEP']}"
            for o in ("AB", "BA")
        ]
        expected.append(row)
    if tables[2][1:] != expected:
        raise ValueError("Seed/order table differs from experiment")
    refs = read(PAPER / "sources.json")["references"]
    manuscript = (PAPER / "manuscript.md").read_text(encoding="utf-8")
    for citation in refs:
        author, year = citation["key"].rsplit(", ", 1)
        pattern = re.escape(author) + r"(?:, | \()" + re.escape(year)
        combined_qwen = (
            citation["key"] == "Qwen Team, 2026b" and "Qwen Team, 2026a, 2026b" in manuscript
        )
        if re.search(pattern, manuscript) is None and not combined_qwen:
            raise ValueError("Uncited reference: " + citation["key"])
    if "et al.." in body:
        raise ValueError("Duplicate reference punctuation")
    if check_files:
        manifest = read(PAPER / "publication.json")
        for name, expected_digest in manifest["files"].items():
            path = (PAPER / name).resolve()
            if not path.is_relative_to(PAPER.resolve()) or digest(path) != expected_digest:
                raise ValueError("Changed manuscript artifact: " + name)
        for name, expected_digest in manifest["inputs"].items():
            path = (ROOT / name).resolve()
            if not path.is_relative_to(ROOT.resolve()) or digest(path) != expected_digest:
                raise ValueError("Changed manuscript input: " + name)
    return {
        "state": "verified",
        "tables": 3,
        "equations": 4,
        "reference_styles_preserved": True,
        "experimental_tables_match": True,
        "references": len(refs),
    }


def seal():
    """Record verified deliverable hashes after native rendering and visual QA."""
    from pypdf import PdfReader

    result = audit(check_files=False)
    names = [
        "manuscript.docx",
        "paper.pdf",
        "manuscript.md",
        "sources.json",
        "style_template.docx",
        "style_reference.json",
        "figures/workflow.png",
        "figures/outcomes.png",
    ]
    inputs = [
        "study/02_confirmation/comparison.json",
        "study/02_confirmation/plan.json",
        "study/02_confirmation/FREEZE.json",
        "paper/publication.json",
        "src/sp_lense/reporting/second_paper.py",
        "src/sp_lense/reporting/render_manuscript.ps1",
    ]
    pdf = PdfReader(PAPER / "paper.pdf")
    metadata = {
        "title": (PAPER / "manuscript.md").read_text(encoding="utf-8").splitlines()[0][2:],
        "experiment_snapshot": "f71d6de346608df374aaa3263ede24823aec66c0",
        "pages": len(pdf.pages),
        "renderer": pdf.metadata.producer,
        "visual_qa": "All page images inspected after native PDF export and Poppler rasterization",
        "files": {name: digest(PAPER / name) for name in names},
        "inputs": {name: digest(ROOT / name) for name in inputs},
        "checks": result,
        "scope": "Complete research manuscript; no claim of peer-review acceptance or submission.",
    }
    (PAPER / "publication.json").write_text(
        json.dumps(metadata, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    return audit()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["build", "audit", "seal"])
    args = parser.parse_args()
    print(json.dumps({"build": build, "audit": audit, "seal": seal}[args.command]()))
